import io
import pdfplumber
import docx
from .llm import generate_json
from .prompts import RESUME_PARSER_SYSTEM

MAX_MB = 5

# Magic byte signatures for file type detection
_PDF_SIG  = b"%PDF"
_DOCX_SIG = b"PK\x03\x04"   # ZIP-based format (OOXML)


class ResumeError(Exception):
    pass


def _detect_file_type(header: bytes) -> str | None:
    """
    Detect file type from the first bytes of the file content (magic bytes).
    Returns 'pdf', 'docx', or None if unrecognised.
    """
    if header[:4] == _PDF_SIG:
        return "pdf"
    if header[:4] == _DOCX_SIG:
        # Both DOCX and plain ZIP start with PK\x03\x04.
        # DOCX files contain [Content_Types].xml as their first entry.
        # We accept any PK-based file here and let python-docx validate it further.
        return "docx"
    return None


def extract_text(file) -> str:
    """
    Extract plain text from an uploaded resume file.

    File-type detection order:
      1. Inspect actual file bytes (magic bytes / file signature) — primary.
      2. Fall back to the file extension if the header is inconclusive.
    """
    if file.size > MAX_MB * 1024 * 1024:
        raise ResumeError(f"File too large (max {MAX_MB} MB).")

    data = file.read()

    # --- Primary: magic-byte detection ---
    file_type = _detect_file_type(data[:8])

    # --- Fallback: extension hint ---
    if file_type is None:
        name_lower = file.name.lower()
        if name_lower.endswith(".pdf"):
            file_type = "pdf"
        elif name_lower.endswith(".docx"):
            file_type = "docx"

    if file_type is None:
        raise ResumeError(
            "Unsupported file type. Please upload a PDF or DOCX resume."
        )

    try:
        if file_type == "pdf":
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        else:  # docx
            d = docx.Document(io.BytesIO(data))
            text = "\n".join(p.text for p in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    text += "\n" + " | ".join(c.text for c in row.cells)
    except Exception:
        raise ResumeError(
            "Could not read this file — it may be corrupted or password-protected."
        )

    text = text.strip()
    if len(text) < 100:
        raise ResumeError(
            "No readable text found. "
            "Please upload a text-based PDF or DOCX (not a scanned image)."
        )
    return text


def parse_resume(text: str) -> dict:
    try:
        data = generate_json(f"Resume text:\n{text[:12000]}", RESUME_PARSER_SYSTEM)
    except Exception:
        raise ResumeError("Could not analyse the resume right now. Please try again.")
    if not isinstance(data, dict) or not data.get("skills"):
        raise ResumeError("This doesn't look like a resume (no skills found).")
    return data